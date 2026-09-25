# -*- coding: utf-8 -*-
"""
Gera as planilhas do kit ROUANET 31/10.

Uso:
    python3 fonte/build_xlsx.py            -> gera versão em branco e versão EXEMPLO em produto/

Todas as regras numéricas ficam na aba PARAMETROS (células amarelas).
Se a norma mudar, o comprador altera ali e todas as calculadoras se atualizam.
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter
from datetime import date

OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "produto")

# ---------------------------------------------------------------- estilos
NAVY = "1F2A44"
ORANGE = "E85D04"
F_HEAD = PatternFill("solid", fgColor=NAVY)
F_SUB = PatternFill("solid", fgColor="DDE3EE")
F_INPUT = PatternFill("solid", fgColor="FFF6D5")
F_CALC = PatternFill("solid", fgColor="EEF1F5")
F_TITLE = PatternFill("solid", fgColor=ORANGE)
F_RED = PatternFill("solid", fgColor="F8D7DA")
F_YEL = PatternFill("solid", fgColor="FFF3CD")
F_GRN = PatternFill("solid", fgColor="D4EDDA")
WHITE_B = Font(color="FFFFFF", bold=True)
BOLD = Font(bold=True)
TITLE_FONT = Font(color="FFFFFF", bold=True, size=16)
THIN = Side(style="thin", color="B8C0CC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
BRL = 'R$ #,##0.00'
PCT = '0.0%'
DATE = 'DD/MM/YYYY'

N_ORC = 200  # linhas de orçamento
ORC_FIRST, ORC_LAST = 8, 8 + N_ORC - 1
N_CRONO = 60
CR_FIRST, CR_LAST = 8, 8 + N_CRONO - 1

ETAPAS = [
    "Pré-produção", "Produção / Execução", "Pós-produção", "Divulgação / Comunicação",
    "Acessibilidade", "Custos administrativos", "Captação de recursos", "Encerramento",
]
CATEGORIAS = [
    # nome, limite unitário (cachês) ou vazio
    ("Produção / atividade-fim", None),
    ("Cachê – artista solo", "L_CACHE_SOLO"),
    ("Cachê – grupo ou coletivo", "L_CACHE_GRUPO"),
    ("Cachê – músico de orquestra", "L_CACHE_MUSICO"),
    ("Cachê – maestro / regente", "L_CACHE_MAESTRO"),
    ("Cachê – palestrante / conferencista", "L_CACHE_PALESTRA"),
    ("Remuneração do proponente", None),
    ("Custos de administração", None),
    ("Divulgação / comunicação", None),
    ("Acessibilidade / comunicação e divulgação acessíveis", None),
    ("Remuneração de captação", None),
    ("Tributos e encargos", None),
    ("Outros (explique na observação)", None),
]
FONTES = ["Incentivo fiscal (Rouanet)", "Outras fontes"]


def title(ws, text, sub=None, width_cols=10):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width_cols)
    c = ws.cell(row=1, column=1, value=text)
    c.font = TITLE_FONT
    c.fill = F_TITLE
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    if sub:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width_cols)
        s = ws.cell(row=2, column=1, value=sub)
        s.alignment = WRAP
        s.font = Font(italic=True, color="44506A")
        ws.row_dimensions[2].height = 45


def header(ws, row, labels, fill=F_HEAD, font=WHITE_B):
    for i, lab in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=lab)
        c.fill = fill
        c.font = font
        c.alignment = CENTER
        c.border = BORDER
    ws.row_dimensions[row].height = 42


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def style_range(ws, r1, r2, c1, c2, fill=None, fmt=None, wrap=True):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            if fill:
                cell.fill = fill
            if fmt:
                cell.number_format = fmt
            if wrap:
                cell.alignment = WRAP


def add_list(ws, rng, options=None, formula=None):
    if formula is None:
        formula = '"' + ",".join(options) + '"'
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.error = "Escolha uma opção da lista."
    ws.add_data_validation(dv)
    dv.add(rng)


def traffic_light(ws, rng, first_cell):
    """Colore células que começam com 🔴 / 🟡 / 🟢."""
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({first_cell},2)="🔴"'], fill=F_RED))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({first_cell},2)="🟡"'], fill=F_YEL))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({first_cell},2)="🟢"'], fill=F_GRN))
    # ISNUMBER(SEARCH()) cobre alertas concatenados (texto começando com outro emoji)
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("🔴",{first_cell}))'], fill=F_RED))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("🟡",{first_cell}))'], fill=F_YEL))


def name(wb, nm, ref):
    dn = DefinedName(nm, attr_text=ref)
    wb.defined_names[nm] = dn


def status_formula(pct, lim):
    return (f'=IF(OR({pct}="",{lim}=""),"",IF({pct}>{lim},"🔴 ACIMA DO LIMITE",'
            f'IF({pct}>={lim}*MARGEM,"🟡 VERIFICAR","🟢 DENTRO DO LIMITE")))')


# ======================================================================
def build(example=False):
    wb = Workbook()

    # ------------------------------------------------------------ COMECE AQUI
    ws = wb.active
    ws.title = "COMECE AQUI"
    title(ws, "ROUANET 31/10 — Planilhas de Execução", "Kit de organização e conferência. Não garante aprovação nem captação. "
          "Base normativa: Lei 8.313/1991, Decreto 11.453/2023 e IN MinC nº 29/2026 (verificação: 25/09/2026). "
          "Sempre confira a norma vigente no site do Ministério da Cultura.", 4)
    rows = [
        ("ORDEM", "ABA", "O QUE FAZER", "TEMPO"),
        ("1", "PARAMETROS", "Confira se os limites estão iguais aos da norma vigente. Só altere se a norma mudar (células amarelas).", "5 min"),
        ("2", "MAPA DO PROJETO", "Preencha a ideia do projeto campo a campo. É o texto-base do Salic.", "2–3 h"),
        ("3", "ORÇAMENTO", "Lance cada despesa. Preencha também as colunas cinza de conferência (categoria, vínculo, preço, dupla cobertura).", "2–4 h"),
        ("4", "CRONOGRAMA", "Lance as atividades com início e fim. O gráfico de barras e os alertas aparecem sozinhos.", "1 h"),
        ("5", "DISTRIBUICAO", "Informe quantos ingressos/produtos existem e como serão distribuídos.", "30 min"),
        ("6", "RADAR DE LIMITES", "Informe tipo de proponente e cenário de captação. Leia os status 🟢🟡🔴.", "20 min"),
        ("7", "MATRIZ CONSISTENCIA", "Ligue cada objetivo a meta, produto, despesa, atividade e público.", "40 min"),
        ("8", "CHECKLIST DOCUMENTOS", "Marque o que se aplica e o que você já tem.", "30 min"),
        ("9", "CHECKLIST SALIC", "Use ao preencher o sistema, etapa por etapa.", "durante o Salic"),
        ("10", "CHECKLIST FINAL", "50 verificações antes de clicar em enviar. Meta: zero 'NÃO'.", "40 min"),
        ("11", "CONTROLE NORMATIVO", "Registre a data em que você conferiu a norma oficial.", "5 min"),
    ]
    for i, r in enumerate(rows, start=4):
        for j, v in enumerate(r, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = BORDER
            c.alignment = WRAP
            if i == 4:
                c.fill = F_HEAD
                c.font = WHITE_B
    ws.cell(row=17, column=1, value="LEGENDA DE CORES").font = BOLD
    leg = [(F_INPUT, "Amarelo claro = você preenche"), (F_CALC, "Cinza = fórmula automática (não apague)"),
           (F_GRN, "🟢 = dentro do limite / conferido / coerente"), (F_YEL, "🟡 = verificar / revisar (perto do limite ou informação faltando)"),
           (F_RED, "🔴 = acima do limite / risco alto / inconsistente")]
    for i, (f, t) in enumerate(leg, start=18):
        ws.cell(row=i, column=1).fill = f
        ws.cell(row=i, column=2, value=t)
    ws.cell(row=24, column=1, value="REGRA DE OURO").font = Font(bold=True, color=ORANGE, size=13)
    ws.merge_cells("B24:D24")
    ws.cell(row=24, column=2, value="SE NÃO CONSIGO EXPLICAR POR QUE ESSA DESPESA EXISTE, NÃO COLOQUE A DESPESA.").font = Font(bold=True, size=13)
    ws.merge_cells("A26:D28")
    ws.cell(row=26, column=1, value=(
        "IMPORTANTE: as calculadoras são ferramentas de CONFERÊNCIA prévia. O Salic aplica as fórmulas oficiais e o Ministério da Cultura "
        "faz a análise técnica. 🟢 na planilha não significa aprovação — significa apenas que o número informado está dentro do limite "
        "configurado na aba PARAMETROS.")).alignment = WRAP
    widths(ws, [10, 26, 90, 16])

    # ------------------------------------------------------------ PARAMETROS
    ws = wb.create_sheet("PARAMETROS")
    title(ws, "PARÂMETROS NORMATIVOS (IN MinC nº 29/2026)", "Células amarelas = limites usados por todas as calculadoras. "
          "Confira cada linha no texto oficial antes de usar. Se a norma mudar, altere aqui.", 7)
    header(ws, 4, ["Regra", "Limite", "Base de cálculo", "Quem está sujeito", "Exceções / observações", "Fonte (confira no texto oficial)", "Status da verificação"])
    params = [
        ("L_REM_GERAL", "Remuneração do proponente – regra geral", 0.20, PCT, "VALOR CAPTADO",
         "Proponente que presta serviço ao projeto previsto no orçamento analítico. Pagamentos a cônjuge, companheiro(a), empresa coligada ou com sócio em comum SOMAM neste limite.",
         "Não se aplica a grupos artísticos familiares, corpos artísticos estáveis e grupos/coletivos culturais que atuem na execução do projeto (ver texto oficial). PF e MEI têm limite próprio (linha abaixo).",
         "IN MinC 29/2026 – dispositivo sobre remuneração do proponente"),
        ("L_REM_PFMEI", "Remuneração do proponente – pessoa física ou MEI", 0.30, PCT, "VALOR CAPTADO",
         "Proponente pessoa física ou microempreendedor individual.", "Continua exigindo serviço efetivo previsto no orçamento.",
         "IN MinC 29/2026 – mesmo dispositivo"),
        ("L_FORNECEDOR", "Pagamento a um mesmo fornecedor", 0.20, PCT, "VALOR CAPTADO",
         "Qualquer fornecedor pago com recurso incentivado.", "Há exceções previstas (ex.: intervenções de conservação e restauro de bens culturais). Confira a lista completa no texto oficial.",
         "IN MinC 29/2026 – dispositivo sobre limite por fornecedor"),
        ("L_ADM", "Custos de administração", 0.15, PCT, "VALOR DO PROJETO",
         "Todos os projetos.", "Pagos proporcionalmente às parcelas captadas. 2026: regularização documental de bens de patrimônio passou a ser admitida como custo administrativo; caiu a regra de não concentrar mais de 50% em uma única despesa administrativa.",
         "IN MinC 29/2026 – custos vinculados"),
        ("L_ACESS", "Acessibilidade, comunicação e divulgação acessíveis", 0.20, PCT, "VALOR DO PROJETO",
         "Todos os projetos.", "Gastos proporcionalmente à captação. Inclui consultor/coordenador de acessibilidade e equipe treinada para atender PcD durante o evento.",
         "IN MinC 29/2026 – custos vinculados"),
        ("L_DIVULG", "Divulgação / comunicação", 0.20, PCT, "VALOR DO PROJETO",
         "Todos os projetos.", "Inclui assessoria de comunicação, despesas de divulgação e impulsionamento de conteúdo. Materiais e aplicação de marcas passam por aprovação do MinC.",
         "IN MinC 29/2026 – dispositivo sobre divulgação"),
        ("L_CAPT_PCT", "Remuneração de captação – percentual", 0.10, PCT, "VALOR DO PROJETO",
         "Quem contrata serviço de captação.", "Paga proporcionalmente ao valor efetivamente captado. Serviço deve ser prestado ao PROPONENTE; é proibido pagar serviço prestado ao incentivador.",
         "IN MinC 29/2026 – remuneração de captação"),
        ("L_CAPT_TETO", "Remuneração de captação – teto em reais", 150000, BRL, "valor absoluto",
         "Idem.", "Vale o MENOR entre 10% do valor do projeto e R$ 150.000,00.", "IN MinC 29/2026 – remuneração de captação"),
        ("L_CACHE_SOLO", "Cachê – artista solo (por apresentação)", 25000, BRL, "por apresentação",
         "Pago com recurso incentivado.", "Valores maiores: só com outras fontes ou mediante apreciação da CNIC.", "IN MinC 29/2026 – limites de cachê"),
        ("L_CACHE_GRUPO", "Cachê – grupo / coletivo (por apresentação)", 50000, BRL, "por apresentação", "Idem.", "Idem.", "IN MinC 29/2026 – limites de cachê"),
        ("L_CACHE_MUSICO", "Cachê – músico de orquestra", 5000, BRL, "por apresentação", "Idem.", "Idem.", "IN MinC 29/2026 – limites de cachê"),
        ("L_CACHE_MAESTRO", "Cachê – maestro / regente", 25000, BRL, "por apresentação", "Idem.", "Idem.", "IN MinC 29/2026 – limites de cachê"),
        ("L_CACHE_PALESTRA", "Cachê – palestrante / conferencista", 5000, BRL, "por participação", "Idem.", "Idem.", "IN MinC 29/2026 – limites de cachê"),
        ("L_INICIANTE", "Primeiro projeto no Pronac – dispensa de comprovar atuação cultural", 200000, BRL, "VALOR DO PROJETO (confira o termo exato)",
         "Proponente que apresenta seu PRIMEIRO projeto ao Pronac.", "Dispensa SOMENTE a comprovação de atuação na área cultural. NÃO dispensa cadastro, documentos, orçamento, acessibilidade, democratização nem análise técnica. Acima do limite: precisa comprovar atuação (portfólio no Salic).",
         "IN MinC 29/2026 – dispositivo sobre comprovação de atuação"),
        ("L_DIST_PATROC", "Distribuição gratuita promocional – patrocinador (máximo)", 0.10, PCT, "total de ingressos/produtos", "Projetos com produtos/ingressos.", "", "IN MinC 29/2026 – plano de distribuição"),
        ("L_DIST_PROPON", "Distribuição gratuita promocional – proponente/divulgação (máximo)", 0.10, PCT, "total de ingressos/produtos", "Idem.", "", "IN MinC 29/2026 – plano de distribuição"),
        ("L_DIST_SOCIAL", "Distribuição gratuita social/educativa (mínimo)", 0.10, PCT, "total de ingressos/produtos", "Idem.", "", "IN MinC 29/2026 – plano de distribuição"),
        ("L_DIST_POPULAR", "Comercialização a preço popular (mínimo)", 0.20, PCT, "total de ingressos/produtos", "Idem.", "", "IN MinC 29/2026 – plano de distribuição"),
        ("L_PRECO_POPULAR", "Preço popular – valor máximo", 50, BRL, "por ingresso/produto", "Idem.", "", "IN MinC 29/2026 – plano de distribuição"),
        ("L_TICKET_MEDIO", "Preço médio do ingresso (máximo) – CONFIRA", 250, BRL, "média dos ingressos comercializados", "Idem.",
         "Valor encontrado em fontes secundárias. CONFIRA no texto oficial antes de usar.", "IN MinC 29/2026 – plano de distribuição"),
        ("L_PF_QTD", "Pessoa física – projetos ativos", 2, '0', "quantidade", "Proponente PF.", "Teto global na linha abaixo.", "IN MinC 29/2026 – limites por proponente"),
        ("L_PF_TETO", "Pessoa física – teto global", 500000, BRL, "soma dos projetos ativos", "Proponente PF.", "", "IN MinC 29/2026 – limites por proponente"),
        ("L_MEI_QTD", "MEI – projetos ativos", 4, '0', "quantidade", "Proponente MEI.", "", "IN MinC 29/2026 – limites por proponente"),
        ("L_MEI_TETO", "MEI – teto global", 1500000, BRL, "soma dos projetos ativos", "Proponente MEI.", "", "IN MinC 29/2026 – limites por proponente"),
        ("L_PJ_QTD", "Demais pessoas jurídicas – projetos ativos", 10, '0', "quantidade", "Demais PJ.", "2026: reduzido de 16 para 10; não há mais limite separado para optantes do Simples.", "IN MinC 29/2026 – limites por proponente"),
        ("L_PJ_TETO", "Demais pessoas jurídicas – teto global", 15000000, BRL, "soma dos projetos ativos", "Demais PJ.", "Existem regras específicas (ex.: planos anuais, patrimônio, regiões). Confira.", "IN MinC 29/2026 – limites por proponente"),
        ("MARGEM", "Margem de alerta amarelo (a partir de X% do limite)", 0.90, PCT, "—", "Configuração do kit (não é regra da IN).", "Ex.: 0,90 → acima de 90% do limite aparece 🟡 VERIFICAR.", "Configuração interna"),
    ]
    for i, (nm, rule, val, fmt, base, who, exc, src) in enumerate(params, start=5):
        vals = [rule, val, base, who, exc, src, "☐ conferido no texto oficial em ___/___/2026"]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = BORDER
            c.alignment = WRAP
        ws.cell(row=i, column=2).fill = F_INPUT
        ws.cell(row=i, column=2).number_format = fmt
        name(wb, nm, f"PARAMETROS!$B${i}")
    last_param = 4 + len(params)
    widths(ws, [44, 16, 22, 34, 60, 34, 26])

    # tabela de categorias (para listas) na coluna I/J
    ws.cell(row=4, column=9, value="Categorias de limite (listas)").font = BOLD
    ws.cell(row=4, column=10, value="Limite unitário").font = BOLD
    for i, (cat, lim) in enumerate(CATEGORIAS, start=5):
        ws.cell(row=i, column=9, value=cat)
        if lim:
            ws.cell(row=i, column=10, value=f"={lim}").number_format = BRL
    cat_last = 4 + len(CATEGORIAS)
    name(wb, "LISTA_CATEGORIAS", f"PARAMETROS!$I$5:$I${cat_last}")
    name(wb, "TAB_CATEGORIAS", f"PARAMETROS!$I$5:$J${cat_last}")
    ws.cell(row=4, column=12, value="Etapas").font = BOLD
    for i, e in enumerate(ETAPAS, start=5):
        ws.cell(row=i, column=12, value=e)
    name(wb, "LISTA_ETAPAS", f"PARAMETROS!$L$5:$L${4 + len(ETAPAS)}")
    ws.column_dimensions["I"].width = 46
    ws.column_dimensions["J"].width = 16
    ws.column_dimensions["L"].width = 28
    ws.freeze_panes = "B5"

    # ------------------------------------------------------------ MAPA DO PROJETO
    ws = wb.create_sheet("MAPA DO PROJETO")
    title(ws, "PLANILHA 1 — MAPA DO PROJETO", "Preencha a coluna C. A coluna D mostra se o campo está preenchido e quantos caracteres tem. "
          "Escreva para o seu projeto — nada de texto genérico copiado.", 5)
    header(ws, 4, ["Campo", "Pergunta que o campo responde", "SUA RESPOSTA", "Status", "Nº de caracteres"])
    campos = [
        ("Nome do projeto", "Como o projeto se chama? Curto, claro, sem siglas soltas."),
        ("Área cultural", "Qual a área (ex.: artes cênicas, música, audiovisual, patrimônio, artes visuais, humanidades…)?"),
        ("Segmento", "Qual o segmento dentro da área? Ele define enquadramento (art. 18 ou art. 26 da Lei 8.313/1991)."),
        ("Resumo", "Em 5 linhas: O QUÊ + ONDE + QUANDO + PARA QUEM + COMO."),
        ("Objeto (o que será entregue)", "Qual é a entrega principal? (ex.: 8 apresentações gratuitas de teatro de rua)."),
        ("Objetivo geral", "Qual a grande mudança/resultado cultural que o projeto busca? (1 frase, verbo no infinitivo)."),
        ("Objetivos específicos", "3 a 5 objetivos mensuráveis. Cada um vira META e PRODUTO."),
        ("Justificativa", "Por que este projeto, neste lugar, agora? Qual a lacuna? Qual a relação com a Lei 8.313/1991 (art. 1º)?"),
        ("Metodologia", "Como será feito, etapa por etapa? Quem faz o quê? Cada etapa precisa ter despesa no orçamento."),
        ("Público-alvo", "Quem é o público (perfil, faixa etária, território)? Quantas pessoas?"),
        ("Beneficiários", "Quem recebe gratuidade, formação, acesso? Escolas, comunidades, instituições?"),
        ("Produtos", "Lista dos produtos culturais (apresentação, livro, oficina, exposição, vídeo…) com quantidades."),
        ("Metas", "Números: quantas apresentações, vagas, exemplares, horas, pessoas atendidas."),
        ("Local de realização", "Município(s)/UF e tipo de espaço. Tem anuência/cessão do espaço?"),
        ("Período de execução", "Data de início e fim previstas (máximo de 36 meses de execução, conforme IN 29/2026)."),
        ("Etapas", "Pré-produção, produção, pós-produção, divulgação, execução, encerramento — o que acontece em cada uma?"),
        ("Acessibilidade (física, comunicacional, atitudinal)", "Quais medidas concretas, em quais produtos, com qual custo?"),
        ("Democratização do acesso", "Gratuidade, preço popular, ações formativas, distribuição — números e onde."),
        ("Resultados esperados", "O que muda depois do projeto? Como vai comprovar (fotos, listas, relatórios, clipping)?"),
        ("Outras fontes de recurso", "Existe dinheiro/apoio além do incentivo? Está confirmado? Qual despesa ele paga?"),
    ]
    for i, (c1, c2) in enumerate(campos, start=5):
        ws.cell(row=i, column=1, value=c1).font = BOLD
        ws.cell(row=i, column=2, value=c2)
        ws.cell(row=i, column=3).fill = F_INPUT
        ws.cell(row=i, column=4, value=f'=IF(LEN(TRIM(C{i}))=0,"🔴 VAZIO",IF(LEN(TRIM(C{i}))<40,"🟡 MUITO CURTO","🟢 PREENCHIDO"))').fill = F_CALC
        ws.cell(row=i, column=5, value=f"=LEN(C{i})").fill = F_CALC
        for col in range(1, 6):
            ws.cell(row=i, column=col).border = BORDER
            ws.cell(row=i, column=col).alignment = WRAP
        ws.row_dimensions[i].height = 70
    last_mapa = 4 + len(campos)
    traffic_light(ws, f"D5:D{last_mapa}", "D5")
    r = last_mapa + 2
    ws.cell(row=r, column=1, value="Campos preenchidos").font = BOLD
    ws.cell(row=r, column=3, value=f'=COUNTIF(D5:D{last_mapa},"🟢*")&" de {len(campos)}"')
    widths(ws, [30, 52, 80, 18, 12])
    ws.freeze_panes = "C5"

    # ------------------------------------------------------------ ORÇAMENTO
    ws = wb.create_sheet("ORÇAMENTO")
    title(ws, "PLANILHA 2 — ORÇAMENTO ANALÍTICO", "Colunas A–J: o que vai para o Salic. Colunas K–O (cinza): SUAS conferências. "
          "Colunas P–Q: alertas automáticos. Regra: se não consegue explicar por que a despesa existe, não coloque.", 17)
    ws["A3"] = "Total incentivo (Rouanet)"; ws["A3"].font = BOLD
    ws["C3"] = f'=SUMIFS(G{ORC_FIRST}:G{ORC_LAST},H{ORC_FIRST}:H{ORC_LAST},"Incentivo*")'
    ws["A4"] = "Total outras fontes"; ws["A4"].font = BOLD
    ws["C4"] = f'=SUMIFS(G{ORC_FIRST}:G{ORC_LAST},H{ORC_FIRST}:H{ORC_LAST},"Outras*")'
    ws["A5"] = "TOTAL GERAL"; ws["A5"].font = BOLD
    ws["C5"] = "=C3+C4"
    for a in ("C3", "C4", "C5"):
        ws[a].number_format = BRL
        ws[a].fill = F_CALC
        ws[a].font = BOLD
    ws["E3"] = "Linhas com alerta 🔴"; ws["E3"].font = BOLD
    ws["G3"] = f'=COUNTIF(Q{ORC_FIRST}:Q{ORC_LAST},"*🔴*")'
    ws["E4"] = "Linhas com alerta 🟡"; ws["E4"].font = BOLD
    ws["G4"] = f'=COUNTIF(Q{ORC_FIRST}:Q{ORC_LAST},"*🟡*")'
    ws["E5"] = "Linhas 🟢 conferidas"; ws["E5"].font = BOLD
    ws["G5"] = f'=COUNTIF(Q{ORC_FIRST}:Q{ORC_LAST},"🟢*")'
    for a in ("G3", "G4", "G5"):
        ws[a].fill = F_CALC
        ws[a].font = BOLD
    name(wb, "VALOR_INCENTIVO", "'ORÇAMENTO'!$C$3")
    name(wb, "VALOR_OUTRAS", "'ORÇAMENTO'!$C$4")
    name(wb, "VALOR_TOTAL", "'ORÇAMENTO'!$C$5")

    header(ws, 7, [
        "Etapa", "Item", "Descrição", "Unidade", "Quantidade", "Valor unitário", "Valor total", "Fonte", "Fornecedor", "Observação",
        "Categoria de limite", "Pessoa vinculada ao proponente? (cônjuge, companheiro, coligada, sócio em comum)",
        "POR QUE ESSA DESPESA EXISTE? (vínculo com a metodologia)", "Preço de mercado conferido?", "Esta despesa é paga também por outra fonte?",
        "% do fornecedor sobre o VALOR CAPTADO", "ALERTAS AUTOMÁTICOS",
    ])
    for col in range(11, 16):
        ws.cell(row=7, column=col).fill = PatternFill("solid", fgColor="44506A")
    for col in (16, 17):
        ws.cell(row=7, column=col).fill = F_TITLE
    rg = lambda c: f"{c}{ORC_FIRST}:{c}{ORC_LAST}"
    for r in range(ORC_FIRST, ORC_LAST + 1):
        ws.cell(row=r, column=7, value=f'=IF(OR(E{r}="",F{r}=""),"",E{r}*F{r})')
        ws.cell(row=r, column=16, value=(
            f'=IF(OR(I{r}="",G{r}="",LEFT(H{r},9)<>"Incentivo",VALOR_CAPTADO=0),"",'
            f'SUMIFS($G${ORC_FIRST}:$G${ORC_LAST},$I${ORC_FIRST}:$I${ORC_LAST},I{r},$H${ORC_FIRST}:$H${ORC_LAST},"Incentivo*")/VALOR_CAPTADO)'))
        ws.cell(row=r, column=17, value=(
            f'=IF(AND(B{r}="",G{r}=""),"",'
            f'IF(OR(E{r}="",F{r}=""),"🟡 Falta quantidade ou valor. ","")'
            f'&IF(H{r}="","🟡 Informe a fonte. ","")'
            f'&IF(K{r}="","🟡 Classifique a categoria de limite. ","")'
            f'&IF(TRIM(M{r})="","🔴 Sem justificativa de vínculo: se não explica, não coloque. ","")'
            f'&IF(N{r}<>"Sim","🟡 Preço de mercado não conferido. ","")'
            f'&IF(O{r}="Sim","🔴 Possível DUPLA COBERTURA: a mesma despesa não pode ser paga duas vezes. ","")'
            f'&IF(AND(L{r}="Sim",K{r}<>"Remuneração do proponente"),"🟡 Pessoa vinculada: este valor SOMA no limite de remuneração do proponente. ","")'
            f'&IF(P{r}="","",IF(P{r}>L_FORNECEDOR,"🔴 Fornecedor acima do limite sobre o valor captado (confira exceções). ",IF(P{r}>=L_FORNECEDOR*MARGEM,"🟡 Fornecedor perto do limite. ","")))'
            f'&IF(AND(LEFT(K{r},5)="Cachê",F{r}<>""),IF(F{r}>IFERROR(VLOOKUP(K{r},TAB_CATEGORIAS,2,FALSE),9E+99),"🔴 Cachê acima do limite: só com outras fontes ou apreciação da CNIC. ",""),"")'
            f'&IF(AND(LEFT(H{r},6)="Outras",K{r}<>""),"🟡 Outras fontes: confirme que o recurso existe e está comprometido. ","")'
            f')'))
        # se nenhum alerta -> verde (feito via fórmula envolvendo a de cima)
        f = ws.cell(row=r, column=17).value[1:]
        ws.cell(row=r, column=17, value=f'=IF(AND(B{r}="",G{r}=""),"",IF({f}="","🟢 CONFERIDO",{f}))')
    style_range(ws, ORC_FIRST, ORC_LAST, 1, 6, F_INPUT)
    style_range(ws, ORC_FIRST, ORC_LAST, 8, 15, F_INPUT)
    style_range(ws, ORC_FIRST, ORC_LAST, 7, 7, F_CALC, BRL)
    style_range(ws, ORC_FIRST, ORC_LAST, 16, 16, F_CALC, PCT)
    style_range(ws, ORC_FIRST, ORC_LAST, 17, 17, F_CALC)
    for r in range(ORC_FIRST, ORC_LAST + 1):
        ws.cell(row=r, column=6).number_format = BRL
    add_list(ws, rg("A"), formula="=LISTA_ETAPAS")
    add_list(ws, rg("H"), FONTES)
    add_list(ws, rg("K"), formula="=LISTA_CATEGORIAS")
    add_list(ws, rg("L"), ["Sim", "Não"])
    add_list(ws, rg("N"), ["Sim", "Não"])
    add_list(ws, rg("O"), ["Sim", "Não"])
    traffic_light(ws, rg("Q"), f"Q{ORC_FIRST}")
    widths(ws, [20, 28, 36, 12, 11, 15, 16, 22, 24, 26, 30, 20, 40, 14, 16, 14, 60])
    ws.freeze_panes = f"C{ORC_FIRST}"

    # ------------------------------------------------------------ CRONOGRAMA
    ws = wb.create_sheet("CRONOGRAMA")
    title(ws, "PLANILHA 3 — CRONOGRAMA", "Não prometa no cronograma aquilo que o orçamento não consegue executar. "
          "Toda etapa com atividade precisa ter despesa no ORÇAMENTO (ou justificar por que não tem custo).", 26)
    ws["A3"] = "Início do projeto"; ws["A3"].font = BOLD
    ws["A4"] = "Fim do projeto"; ws["A4"].font = BOLD
    ws["A5"] = "Duração total (meses)"; ws["A5"].font = BOLD
    for a in ("B3", "B4"):
        ws[a].fill = F_INPUT
        ws[a].number_format = DATE
        ws[a].border = BORDER
    ws["B5"] = '=IF(OR(B3="",B4=""),"",DATEDIF(B3,B4,"m")+1)'
    ws["B5"].fill = F_CALC
    ws["C5"] = '=IF(B5="","",IF(B5>36,"🔴 Acima de 36 meses de execução (confira a IN)","🟢 Dentro de 36 meses"))'
    traffic_light(ws, "C5", "C5")
    header(ws, 7, ["Atividade", "Etapa", "Início", "Fim", "Duração (dias)", "Responsável", "Observações", "ALERTAS"])
    ws.cell(row=7, column=8).fill = F_TITLE
    for k in range(18):
        col = 9 + k
        c = ws.cell(row=7, column=col)
        c.value = '=IF($B$3="","",DATE(YEAR($B$3),MONTH($B$3),1))' if k == 0 else f'=IF({get_column_letter(col - 1)}7="","",EDATE({get_column_letter(col - 1)}7,1))'
        c.number_format = "MMM/YY"
        c.fill = F_HEAD
        c.font = WHITE_B
        c.alignment = CENTER
        ws.column_dimensions[get_column_letter(col)].width = 7
    for r in range(CR_FIRST, CR_LAST + 1):
        ws.cell(row=r, column=5, value=f'=IF(OR(C{r}="",D{r}=""),"",D{r}-C{r}+1)')
        ws.cell(row=r, column=8, value=(
            f'=IF(A{r}="","",IF(OR(C{r}="",D{r}=""),"🟡 Falta data. ",'
            f'IF(D{r}<C{r},"🔴 Fim antes do início. ","")'
            f'&IF(OR(AND($B$3<>"",C{r}<$B$3),AND($B$4<>"",D{r}>$B$4)),"🔴 Fora do período do projeto. ",""))'
            f'&IF(B{r}="","🟡 Informe a etapa. ",IF(COUNTIF(\'ORÇAMENTO\'!$A${ORC_FIRST}:$A${ORC_LAST},B{r})=0,"🟡 Etapa sem despesa no orçamento: quem paga? ",""))'
            f'&IF(F{r}="","🟡 Sem responsável. ",""))'))
        f = ws.cell(row=r, column=8).value[1:]
        ws.cell(row=r, column=8, value=f'=IF(A{r}="","",IF({f}="","🟢 OK",{f}))')
        for k in range(18):
            col = 9 + k
            L = get_column_letter(col)
            ws.cell(row=r, column=col, value=f'=IF(OR($C{r}="",$D{r}="",{L}$7=""),"",IF(AND($C{r}<=EOMONTH({L}$7,0),$D{r}>={L}$7),"■",""))').alignment = CENTER
    style_range(ws, CR_FIRST, CR_LAST, 1, 7, F_INPUT)
    for r in range(CR_FIRST, CR_LAST + 1):
        ws.cell(row=r, column=3).number_format = DATE
        ws.cell(row=r, column=4).number_format = DATE
        ws.cell(row=r, column=5).fill = F_CALC
        ws.cell(row=r, column=8).fill = F_CALC
        ws.cell(row=r, column=8).alignment = WRAP
        ws.cell(row=r, column=8).border = BORDER
    add_list(ws, f"B{CR_FIRST}:B{CR_LAST}", formula="=LISTA_ETAPAS")
    traffic_light(ws, f"H{CR_FIRST}:H{CR_LAST}", f"H{CR_FIRST}")
    ws.conditional_formatting.add(f"I{CR_FIRST}:Z{CR_LAST}", FormulaRule(formula=[f'I{CR_FIRST}="■"'], fill=PatternFill("solid", fgColor=ORANGE), font=Font(color=ORANGE)))
    for i, w in enumerate([34, 22, 12, 12, 10, 18, 26, 44], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"B{CR_FIRST}"

    # ------------------------------------------------------------ DISTRIBUIÇÃO
    ws = wb.create_sheet("DISTRIBUICAO")
    title(ws, "CALCULADORA F — PLANO DE DISTRIBUIÇÃO E DEMOCRATIZAÇÃO", "Para projetos com ingressos ou produtos (livros, CDs, exemplares). "
          "Se o seu projeto é 100% gratuito, informe tudo em 'gratuita social/educativa' e 'demais gratuitos' e registre isso no Salic.", 8)
    ws["A4"] = "Total de ingressos / produtos previstos"; ws["A4"].font = BOLD
    ws["C4"].fill = F_INPUT
    header(ws, 6, ["Faixa", "Quantidade", "VALOR INFORMADO (%)", "BASE DE CÁLCULO", "LIMITE APLICÁVEL", "Tipo de limite", "STATUS", "OBSERVAÇÃO"])
    dist = [
        ("Gratuita promocional – patrocinador", "L_DIST_PATROC", "máximo", "Até o limite; é promocional, não conta como democratização."),
        ("Gratuita promocional – proponente (divulgação)", "L_DIST_PROPON", "máximo", "Até o limite; use em ações de divulgação."),
        ("Gratuita com caráter social/educativo", "L_DIST_SOCIAL", "mínimo", "Diga PARA QUEM (escolas, ONGs, comunidade) no Salic."),
        ("Comercialização a preço popular", "L_DIST_POPULAR", "mínimo", "Preço unitário não pode passar do teto de preço popular."),
        ("Demais (comercialização ou outras gratuidades)", None, "—", "Restante. Confira o preço médio abaixo."),
    ]
    for i, (lab, lim, tipo, obs) in enumerate(dist, start=7):
        ws.cell(row=i, column=1, value=lab)
        ws.cell(row=i, column=2).fill = F_INPUT
        ws.cell(row=i, column=3, value=f'=IF(OR($C$4="",$C$4=0,B{i}=""),"",B{i}/$C$4)')
        ws.cell(row=i, column=4, value="Total de ingressos/produtos")
        if lim:
            ws.cell(row=i, column=5, value=f"={lim}")
            if tipo == "máximo":
                st = f'=IF(C{i}="","",IF(C{i}>E{i},"🔴 ACIMA DO LIMITE",IF(C{i}>=E{i}*MARGEM,"🟡 VERIFICAR","🟢 DENTRO DO LIMITE")))'
            elif "preço popular" in lab:
                # sem nenhuma venda prevista, o mínimo de preço popular precisa ser confirmado (projeto 100% gratuito)
                st = (f'=IF(C{i}="","🟡 VERIFICAR — informe a quantidade",IF(AND(SUM($B$10:$B$11)=0,$B$9=$C$4),'
                      f'"🟡 VERIFICAR — projeto 100% gratuito: confirme como declarar no Salic",IF(C{i}<E{i},"🔴 ABAIXO DO MÍNIMO","🟢 DENTRO DO LIMITE")))')
            else:
                st = f'=IF(C{i}="","🟡 VERIFICAR — informe a quantidade",IF(C{i}<E{i},"🔴 ABAIXO DO MÍNIMO","🟢 DENTRO DO LIMITE"))'
            ws.cell(row=i, column=7, value=st)
        ws.cell(row=i, column=6, value=tipo)
        ws.cell(row=i, column=8, value=obs)
        for col in range(1, 9):
            ws.cell(row=i, column=col).border = BORDER
            ws.cell(row=i, column=col).alignment = WRAP
        ws.cell(row=i, column=3).number_format = PCT
        ws.cell(row=i, column=5).number_format = PCT
        ws.cell(row=i, column=3).fill = F_CALC
        ws.cell(row=i, column=7).fill = F_CALC
    ws["A12"] = "Soma das faixas"; ws["A12"].font = BOLD
    ws["B12"] = "=SUM(B7:B11)"
    ws["C12"] = '=IF(OR(C4="",C4=0),"",IF(B12=C4,"🟢 Soma confere com o total","🔴 Soma das faixas ≠ total informado"))'
    ws["A14"] = "Preço do ingresso popular (R$)"; ws["A14"].font = BOLD
    ws["B14"].fill = F_INPUT; ws["B14"].number_format = BRL
    ws["C14"] = '=IF(B14="","",IF(B14>L_PRECO_POPULAR,"🔴 Acima do teto de preço popular","🟢 DENTRO DO LIMITE"))'
    ws["A15"] = "Receita total prevista com vendas (R$)"; ws["A15"].font = BOLD
    ws["B15"].fill = F_INPUT; ws["B15"].number_format = BRL
    ws["A16"] = "Quantidade total comercializada"; ws["A16"].font = BOLD
    ws["B16"].fill = F_INPUT
    ws["A17"] = "Preço médio (R$)"; ws["A17"].font = BOLD
    ws["B17"] = '=IF(OR(B16="",B16=0),"",B15/B16)'; ws["B17"].number_format = BRL; ws["B17"].fill = F_CALC
    ws["C17"] = '=IF(B17="","",IF(B17>L_TICKET_MEDIO,"🔴 Preço médio acima do teto configurado (confira a IN)","🟢 DENTRO DO LIMITE"))'
    ws["A19"] = "Medida adicional de ampliação de acesso prevista? (descreva)"; ws["A19"].font = BOLD
    ws.merge_cells("B19:H19"); ws["B19"].fill = F_INPUT
    ws["A20"] = "Ações formativas / de acesso vinculadas (descreva)"; ws["A20"].font = BOLD
    ws.merge_cells("B20:H20"); ws["B20"].fill = F_INPUT
    traffic_light(ws, "G7:G11", "G7")
    traffic_light(ws, "C12:C17", "C12")
    widths(ws, [46, 14, 16, 26, 14, 12, 26, 50])
    name(wb, "DIST_STATUS1", "DISTRIBUICAO!$G$7")

    # ------------------------------------------------------------ RADAR
    ws = wb.create_sheet("RADAR DE LIMITES")
    title(ws, "PLANILHA 4 — RADAR DE LIMITES (CALCULADORAS A–H)", "Cada linha usa a BASE DE CÁLCULO prevista na IN 29/2026: "
          "remuneração do proponente e fornecedor → VALOR CAPTADO; administração, divulgação, acessibilidade e captação → VALOR DO PROJETO. "
          "Resultado é conferência prévia; o Salic e a análise do MinC prevalecem.", 8)
    inputs = [
        ("Tipo de proponente", "", None, "Escolha: Pessoa física / MEI / Demais pessoa jurídica"),
        ("A remuneração se enquadra em exceção? (grupo artístico familiar, corpo artístico estável, grupo/coletivo que atua na execução)", "Não", None, "Se Sim, a planilha marca 🟡 VERIFICAR: confirme o enquadramento no texto oficial."),
        ("VALOR DO PROJETO (calculado: soma das linhas 'Incentivo')", "=VALOR_INCENTIVO", BRL, "Soma automática do orçamento com fonte Incentivo fiscal."),
        ("VALOR DO PROJETO exibido no Salic (opcional — sobrescreve)", "", BRL, "Se o Salic mostrar valor diferente (ex.: após cálculo dos custos vinculados), digite aqui."),
        ("VALOR DO PROJETO usado nas contas", '=IF(AND(B8<>"",B8>0),B8,B7)', BRL, "Base para administração, divulgação, acessibilidade e captação."),
        ("VALOR TOTAL DO PROJETO (todas as fontes)", "=VALOR_TOTAL", BRL, "Incentivo + outras fontes. Não é base dos limites percentuais do kit."),
        ("Cenário de captação (% do valor do projeto que você espera captar)", 1, PCT, "Teste 100%, 75% e 50%. Remuneração e fornecedor são medidos sobre o que for CAPTADO."),
        ("VALOR CAPTADO (cenário)", "=B9*B11", BRL, "Base para remuneração do proponente e limite por fornecedor."),
        ("É o seu PRIMEIRO projeto apresentado ao Pronac?", "", None, "Sim / Não"),
        ("Nº de projetos ativos que você já tem (sem contar este)", 0, '0', "Consulte seus projetos no Salic."),
        ("Soma dos valores dos projetos ativos (sem contar este)", 0, BRL, "Para conferir o teto global."),
        ("Nº estimado de beneficiários / público total", "", '#,##0', "Para a calculadora H (valor por beneficiário)."),
    ]
    header(ws, 4, ["ENTRADAS", "VALOR", "", "COMO PREENCHER"])
    for i, (lab, val, fmt, how) in enumerate(inputs, start=5):
        ws.cell(row=i, column=1, value=lab).alignment = WRAP
        c = ws.cell(row=i, column=2, value=val if val != "" else None)
        c.fill = F_CALC if isinstance(val, str) and val.startswith("=") else F_INPUT
        if fmt:
            c.number_format = fmt
        ws.cell(row=i, column=4, value=how).alignment = WRAP
        for col in (1, 2, 4):
            ws.cell(row=i, column=col).border = BORDER
        ws.row_dimensions[i].height = 32
    add_list(ws, "B5", ["Pessoa física", "MEI", "Demais pessoa jurídica"])
    add_list(ws, "B6", ["Sim", "Não"])
    add_list(ws, "B13", ["Sim", "Não"])
    name(wb, "VALOR_PROJETO", "'RADAR DE LIMITES'!$B$9")
    name(wb, "VALOR_CAPTADO", "'RADAR DE LIMITES'!$B$12")
    name(wb, "TIPO_PROP", "'RADAR DE LIMITES'!$B$5")

    R0 = 19
    header(ws, R0, ["CALCULADORA", "VALOR INFORMADO", "BASE DE CÁLCULO", "BASE (R$)", "LIMITE APLICÁVEL", "PERCENTUAL RESULTANTE", "STATUS", "OBSERVAÇÃO"])
    O = f"'ORÇAMENTO'!"
    G_, H_, K_, L_ = (f"{O}$G${ORC_FIRST}:$G${ORC_LAST}", f"{O}$H${ORC_FIRST}:$H${ORC_LAST}",
                      f"{O}$K${ORC_FIRST}:$K${ORC_LAST}", f"{O}$L${ORC_FIRST}:$L${ORC_LAST}")
    sumcat = lambda cat: f'SUMIFS({G_},{K_},"{cat}",{H_},"Incentivo*")'
    calc = [
        ("A. Remuneração do proponente (inclui cônjuge/companheiro, coligada, sócio em comum)",
         f'={sumcat("Remuneração do proponente")}+SUMIFS({G_},{L_},"Sim",{K_},"<>Remuneração do proponente",{H_},"Incentivo*")',
         "VALOR CAPTADO", "=VALOR_CAPTADO",
         '=IF(OR(TIPO_PROP="Pessoa física",TIPO_PROP="MEI"),L_REM_PFMEI,L_REM_GERAL)',
         "pct", "rem",
         "Proponente só recebe se PRESTAR serviço previsto no orçamento analítico. Se captar menos, o limite em R$ diminui."),
        ("B. Maior fornecedor (mesmo fornecedor, soma de todas as linhas)",
         f"=IFERROR(MAX({O}$P${ORC_FIRST}:$P${ORC_LAST})*VALOR_CAPTADO,0)", "VALOR CAPTADO", "=VALOR_CAPTADO", "=L_FORNECEDOR", "pct", "std",
         "Veja a coluna P do ORÇAMENTO para saber qual fornecedor. Há exceções na IN (ex.: restauro de bens culturais)."),
        ("C. Custos de administração", f"={sumcat('Custos de administração')}", "VALOR DO PROJETO", "=VALOR_PROJETO", "=L_ADM", "pct", "std",
         "Contador, tarifas, material de escritório do projeto etc. Não disfarce despesa de produção como administração (nem o contrário)."),
        ("D1. Acessibilidade, comunicação e divulgação acessíveis", f"={sumcat('Acessibilidade / comunicação e divulgação acessíveis')}",
         "VALOR DO PROJETO", "=VALOR_PROJETO", "=L_ACESS", "pct", "std",
         "Limite é TETO, não meta. Acessibilidade é obrigatória: 0% aqui é sinal de alerta (ver status)."),
        ("D2. Divulgação / comunicação", f"={sumcat('Divulgação / comunicação')}", "VALOR DO PROJETO", "=VALOR_PROJETO", "=L_DIVULG", "pct", "std",
         "Assessoria de comunicação, peças, impulsionamento. Precisa ter relação com o público do projeto."),
        ("E. Remuneração de captação", f"={sumcat('Remuneração de captação')}", "VALOR DO PROJETO", "=VALOR_PROJETO",
         "=MIN(L_CAPT_PCT*VALOR_PROJETO,L_CAPT_TETO)", "brl", "capt",
         "Limite = MENOR entre 10% do valor do projeto e R$ 150.000. Pago proporcionalmente ao captado. Serviço ao proponente, nunca ao incentivador."),
    ]
    for i, (lab, val, basen, base, lim, kind, mode, obs) in enumerate(calc, start=R0 + 1):
        ws.cell(row=i, column=1, value=lab)
        ws.cell(row=i, column=2, value=val).number_format = BRL
        ws.cell(row=i, column=3, value=basen)
        ws.cell(row=i, column=4, value=base).number_format = BRL
        ws.cell(row=i, column=5, value=lim).number_format = PCT if kind == "pct" else BRL
        ws.cell(row=i, column=6, value=f'=IF(OR(D{i}="",D{i}=0),"",B{i}/D{i})').number_format = PCT
        if mode == "rem":
            st = (f'=IF(F{i}="","",IF(\'RADAR DE LIMITES\'!$B$6="Sim","🟡 VERIFICAR — exceção declarada: confirme o enquadramento",'
                  f'IF(TIPO_PROP="","🟡 VERIFICAR — informe o tipo de proponente",{status_formula(f"F{i}", f"E{i}")[1:]})))')
        elif mode == "capt":
            st = (f'=IF(B{i}=0,"🟢 DENTRO DO LIMITE (sem captador)",IF(B{i}>E{i},"🔴 ACIMA DO LIMITE",'
                  f'IF(B{i}>=E{i}*MARGEM,"🟡 VERIFICAR","🟢 DENTRO DO LIMITE")))')
        else:
            st = status_formula(f"F{i}", f"E{i}")
        if lab.startswith("D1"):
            st = f'=IF(B{i}=0,"🔴 SEM CUSTO DE ACESSIBILIDADE — confira se as medidas estão previstas",{st[1:]})'
        ws.cell(row=i, column=7, value=st)
        ws.cell(row=i, column=8, value=obs)
        for col in range(1, 9):
            ws.cell(row=i, column=col).border = BORDER
            ws.cell(row=i, column=col).alignment = WRAP
            if col in (2, 4, 6, 7):
                ws.cell(row=i, column=col).fill = F_CALC
        ws.row_dimensions[i].height = 48
    r = R0 + len(calc) + 1
    # F distribuição
    ws.cell(row=r, column=1, value="F. Distribuição / democratização (detalhe na aba DISTRIBUICAO)")
    ws.cell(row=r, column=3, value="Total de ingressos/produtos")
    ws.cell(row=r, column=7, value=('=IF(COUNTIF(DISTRIBUICAO!G7:G10,"🔴*")>0,"🔴 ACIMA/ABAIXO DO LIMITE — veja a aba",'
                                    'IF(COUNTIF(DISTRIBUICAO!G7:G10,"🟡*")>0,"🟡 VERIFICAR — veja a aba","🟢 DENTRO DO LIMITE"))'))
    ws.cell(row=r, column=8, value="Faixas máximas (promocionais) e mínimas (social/educativa e preço popular).")
    r += 1
    # H valor por beneficiário
    ws.cell(row=r, column=1, value="H. Valor por beneficiário")
    ws.cell(row=r, column=2, value="=VALOR_TOTAL").number_format = BRL
    ws.cell(row=r, column=3, value="Nº de beneficiários (B16)")
    ws.cell(row=r, column=4, value="=B16").number_format = '#,##0'
    ws.cell(row=r, column=5, value="sem limite normativo")
    ws.cell(row=r, column=6, value='=IF(OR(B16="",B16=0),"",VALOR_TOTAL/B16)').number_format = BRL
    ws.cell(row=r, column=7, value='=IF(F{0}="","🟡 VERIFICAR — informe beneficiários","🟢 Calculado — avalie se é razoável e explique na justificativa")'.format(r))
    ws.cell(row=r, column=8, value="Não há teto na IN para este indicador. Serve para você testar a coerência: custo alto por pessoa precisa de justificativa (ex.: formação intensiva, produto permanente).")
    r += 1
    # I iniciante
    ws.cell(row=r, column=1, value="I. Primeiro projeto: dispensa de comprovar atuação cultural")
    ws.cell(row=r, column=2, value="=MAX(VALOR_PROJETO,VALOR_TOTAL)").number_format = BRL
    ws.cell(row=r, column=3, value="Valor do projeto (o kit usa o MAIOR entre valor do projeto e valor total, por prudência)")
    ws.cell(row=r, column=5, value="=L_INICIANTE").number_format = BRL
    ws.cell(row=r, column=7, value=(f'=IF(B13="","🟡 VERIFICAR — responda se é o primeiro projeto",IF(B13="Não","🟡 Não se aplica: comprove atuação (portfólio no Salic)",'
                                    f'IF(B{r}<=E{r},"🟢 Pode estar dispensado de comprovar atuação — confira a redação oficial","🔴 Acima do limite: precisa comprovar atuação cultural")))'))
    ws.cell(row=r, column=8, value="A dispensa é SÓ da comprovação de atuação na área cultural. Todos os outros requisitos e documentos continuam valendo.")
    r += 1
    # J carteira
    ws.cell(row=r, column=1, value="J. Limite de projetos ativos e teto global do proponente")
    ws.cell(row=r, column=2, value="=B15+VALOR_PROJETO").number_format = BRL
    ws.cell(row=r, column=3, value="Soma dos projetos ativos + este")
    ws.cell(row=r, column=5, value='=IF(TIPO_PROP="Pessoa física",L_PF_TETO,IF(TIPO_PROP="MEI",L_MEI_TETO,IF(TIPO_PROP="","",L_PJ_TETO)))').number_format = BRL
    ws.cell(row=r, column=4, value='=IF(TIPO_PROP="Pessoa física",L_PF_QTD,IF(TIPO_PROP="MEI",L_MEI_QTD,IF(TIPO_PROP="","",L_PJ_QTD)))')
    ws.cell(row=r, column=6, value=f'=IF(OR(E{r}="",E{r}=0),"",B{r}/E{r})').number_format = PCT
    ws.cell(row=r, column=7, value=(f'=IF(TIPO_PROP="","🟡 VERIFICAR — informe o tipo de proponente",IF(OR(B14+1>D{r},B{r}>E{r}),"🔴 ACIMA DO LIMITE",'
                                    f'IF(B{r}>=E{r}*MARGEM,"🟡 VERIFICAR","🟢 DENTRO DO LIMITE")))'))
    ws.cell(row=r, column=8, value="Coluna D mostra a QUANTIDADE máxima de projetos ativos; coluna E o teto global em R$.")
    for rr in range(R0 + len(calc) + 1, r + 1):
        for col in range(1, 9):
            ws.cell(row=rr, column=col).border = BORDER
            ws.cell(row=rr, column=col).alignment = WRAP
        ws.row_dimensions[rr].height = 48
    traffic_light(ws, f"G{R0 + 1}:G{r}", f"G{R0 + 1}")

    # G percentual por categoria
    r += 2
    ws.cell(row=r, column=1, value="G. PERCENTUAL DE CADA CATEGORIA (sobre o VALOR DO PROJETO)").font = Font(bold=True, color=ORANGE, size=12)
    r += 1
    header(ws, r, ["Categoria", "Valor (incentivo)", "% do valor do projeto", "Limite (se houver)"])
    lim_map = {"Remuneração do proponente": "sobre VALOR CAPTADO (ver A)", "Custos de administração": "=L_ADM",
               "Divulgação / comunicação": "=L_DIVULG", "Acessibilidade / comunicação e divulgação acessíveis": "=L_ACESS",
               "Remuneração de captação": "10% / R$150 mil (ver E)"}
    g0 = r + 1
    for i, (cat, lim) in enumerate(CATEGORIAS, start=g0):
        ws.cell(row=i, column=1, value=cat)
        ws.cell(row=i, column=2, value=f"={sumcat(cat)}").number_format = BRL
        ws.cell(row=i, column=3, value=f'=IF(VALOR_PROJETO=0,"",B{i}/VALOR_PROJETO)').number_format = PCT
        v = lim_map.get(cat, "por apresentação (ver PARAMETROS)" if lim else "—")
        c = ws.cell(row=i, column=4, value=v)
        if isinstance(v, str) and v.startswith("="):
            c.number_format = PCT
        for col in range(1, 5):
            ws.cell(row=i, column=col).border = BORDER
    gl = g0 + len(CATEGORIAS)
    ws.cell(row=gl, column=1, value="TOTAL").font = BOLD
    ws.cell(row=gl, column=2, value=f"=SUM(B{g0}:B{gl - 1})").number_format = BRL
    ws.cell(row=gl, column=3, value=f'=IF(VALOR_PROJETO=0,"",B{gl}/VALOR_PROJETO)').number_format = PCT
    ws.cell(row=gl + 1, column=1, value="Linhas do orçamento sem categoria (incentivo)")
    ws.cell(row=gl + 1, column=2, value=f'=SUMIFS({G_},{K_},"",{H_},"Incentivo*")').number_format = BRL
    ws.cell(row=gl + 1, column=3, value=f'=IF(B{gl + 1}>0,"🟡 Classifique todas as linhas","🟢 Tudo classificado")')
    traffic_light(ws, f"C{gl + 1}", f"C{gl + 1}")

    # Cenários de captação
    r = gl + 3
    ws.cell(row=r, column=1, value="E SE EU CAPTAR MENOS? — remuneração do proponente por cenário").font = Font(bold=True, color=ORANGE, size=12)
    r += 1
    header(ws, r, ["Cenário de captação", "Valor captado", "Limite de remuneração em R$", "Remuneração prevista", "% sobre o captado", "Status"])
    rem_row = R0 + 1
    for k, p in enumerate([1, 0.75, 0.5, 0.25], start=r + 1):
        ws.cell(row=k, column=1, value=p).number_format = '0%'
        ws.cell(row=k, column=2, value=f"=VALOR_PROJETO*A{k}").number_format = BRL
        ws.cell(row=k, column=3, value=f"=B{k}*$E${rem_row}").number_format = BRL
        ws.cell(row=k, column=4, value=f"=$B${rem_row}").number_format = BRL
        ws.cell(row=k, column=5, value=f'=IF(B{k}=0,"",D{k}/B{k})').number_format = PCT
        ws.cell(row=k, column=6, value=f'=IF(E{k}="","",IF(E{k}>$E${rem_row},"🔴 Reduza a remuneração neste cenário","🟢 OK"))')
        for col in range(1, 7):
            ws.cell(row=k, column=col).border = BORDER
    traffic_light(ws, f"F{r + 1}:F{r + 4}", f"F{r + 1}")
    ws.cell(row=r + 5, column=1, value=("Leitura: a remuneração do proponente é limitada pelo VALOR CAPTADO. Se o projeto captar só parte do valor, "
                                        "o valor em reais que você pode receber cai junto. Planeje a execução parcial desde já.")).alignment = WRAP
    ws.merge_cells(start_row=r + 5, start_column=1, end_row=r + 5, end_column=6)
    ws.row_dimensions[r + 5].height = 36
    widths(ws, [52, 18, 26, 18, 20, 16, 44, 60])
    ws.freeze_panes = "A5"

    # ------------------------------------------------------------ MATRIZ
    ws = wb.create_sheet("MATRIZ CONSISTENCIA")
    title(ws, "PLANILHA 7 — MATRIZ DE CONSISTÊNCIA", "Parte 1: cada objetivo específico precisa virar meta → produto → despesa → atividade → público. "
          "Copie o NOME EXATO do item do orçamento e da atividade do cronograma: a planilha confere se eles existem.", 10)
    header(ws, 4, ["Objetivo específico", "Meta (com número)", "Produto", "Item do ORÇAMENTO (nome exato)", "Atividade do CRONOGRAMA (nome exato)",
                   "Público / faixa do plano de distribuição", "Item existe no orçamento?", "Atividade existe no cronograma?", "Revisão manual", "RESULTADO"])
    for r in range(5, 20):
        ws.cell(row=r, column=7, value=f'=IF(D{r}="","",IF(COUNTIF(\'ORÇAMENTO\'!$B${ORC_FIRST}:$B${ORC_LAST},D{r})>0,"Sim","Não"))')
        ws.cell(row=r, column=8, value=f'=IF(E{r}="","",IF(COUNTIF(CRONOGRAMA!$A${CR_FIRST}:$A${CR_LAST},E{r})>0,"Sim","Não"))')
        ws.cell(row=r, column=10, value=(
            f'=IF(A{r}="","",IF(OR(COUNTBLANK(B{r}:F{r})>2,G{r}="Não",H{r}="Não"),"🔴 INCONSISTENTE",'
            f'IF(OR(COUNTBLANK(B{r}:F{r})>0,I{r}="Revisar"),"🟡 REVISAR","🟢 COERENTE")))'))
    style_range(ws, 5, 19, 1, 6, F_INPUT)
    style_range(ws, 5, 19, 7, 8, F_CALC)
    style_range(ws, 5, 19, 9, 9, F_INPUT)
    style_range(ws, 5, 19, 10, 10, F_CALC)
    add_list(ws, "I5:I19", ["OK", "Revisar"])
    traffic_light(ws, "J5:J19", "J5")
    r = 22
    ws.cell(row=r, column=1, value="PARTE 2 — AS 6 PERGUNTAS DE CONSISTÊNCIA").font = Font(bold=True, color=ORANGE, size=12)
    header(ws, r + 1, ["Verificação", "Pergunta de controle", "Resposta (Sim / Parcial / Não)", "RESULTADO", "Onde corrigir"])
    perguntas = [
        ("Objetivo x Meta", "Todo objetivo específico tem pelo menos uma meta com número?", "MAPA DO PROJETO (objetivos/metas)"),
        ("Meta x Produto", "Toda meta corresponde a um produto cadastrado no plano de distribuição?", "Salic – Plano de distribuição"),
        ("Produto x Orçamento", "Todo produto tem as despesas necessárias para existir (e nenhuma despesa sobra)?", "ORÇAMENTO"),
        ("Atividade x Cronograma", "Toda atividade da metodologia aparece no cronograma com data e responsável?", "CRONOGRAMA"),
        ("Público x Distribuição", "O público descrito bate com as faixas de distribuição (gratuidade, preço popular, locais)?", "DISTRIBUICAO / Salic"),
        ("Despesa x Metodologia", "Toda despesa tem 'por que existe' preenchido e ligado a uma etapa da metodologia?", "ORÇAMENTO coluna M"),
    ]
    for i, (a, b, onde) in enumerate(perguntas, start=r + 2):
        ws.cell(row=i, column=1, value=a).font = BOLD
        ws.cell(row=i, column=2, value=b)
        ws.cell(row=i, column=3).fill = F_INPUT
        ws.cell(row=i, column=4, value=f'=IF(C{i}="","",IF(C{i}="Sim","🟢 COERENTE",IF(C{i}="Parcial","🟡 REVISAR","🔴 INCONSISTENTE")))').fill = F_CALC
        ws.cell(row=i, column=5, value=onde)
        for col in range(1, 6):
            ws.cell(row=i, column=col).border = BORDER
            ws.cell(row=i, column=col).alignment = WRAP
        ws.row_dimensions[i].height = 36
    add_list(ws, f"C{r + 2}:C{r + 7}", ["Sim", "Parcial", "Não"])
    traffic_light(ws, f"D{r + 2}:D{r + 7}", f"D{r + 2}")
    # verificação automática despesa x metodologia
    i = r + 9
    ws.cell(row=i, column=1, value="Checagem automática").font = BOLD
    ws.cell(row=i, column=2, value="Despesas sem 'por que existe' no ORÇAMENTO:")
    ws.cell(row=i, column=4, value=f'=COUNTIFS(\'ORÇAMENTO\'!$G${ORC_FIRST}:$G${ORC_LAST},">0",\'ORÇAMENTO\'!$M${ORC_FIRST}:$M${ORC_LAST},"")')
    ws.cell(row=i + 1, column=2, value="Atividades do cronograma em etapa sem despesa:")
    ws.cell(row=i + 1, column=4, value=f'=COUNTIF(CRONOGRAMA!$H${CR_FIRST}:$H${CR_LAST},"*Etapa sem despesa*")')
    widths(ws, [30, 42, 26, 30, 30, 30, 14, 14, 12, 20])
    ws.freeze_panes = "B5"

    # ------------------------------------------------------------ CHECKLIST SALIC
    ws = wb.create_sheet("CHECKLIST SALIC")
    title(ws, "PLANILHA 5 — CHECKLIST SALIC (DO LOGIN AO ENVIO)", "Use enquanto preenche o sistema. Os nomes das telas podem mudar: "
          "siga sempre o Manual do Proponente vigente (IN 2026) publicado pelo MinC.", 6)
    header(ws, 4, ["Passo", "Item", "Obrigatório?", "Preparado?", "Revisado?", "Observação"])
    salic = [
        ("1 Acesso", "Conta gov.br ativa, com senha e nível de acesso exigido pelo sistema", "Sim"),
        ("1 Acesso", "Acesso ao Salic realizado com o CPF do responsável", "Sim"),
        ("2 Cadastro", "Cadastro do proponente (PF ou PJ) completo e atualizado", "Sim"),
        ("2 Cadastro", "Dados de contato (e-mail e telefone) que você realmente acompanha", "Sim"),
        ("2 Cadastro", "Vinculação do responsável/procurador ao proponente PJ, quando aplicável", "Depende"),
        ("2 Cadastro", "Portfólio/comprovação de atuação cultural cadastrado no Salic (ou confirmação da dispensa de iniciante)", "Depende"),
        ("3 Nova proposta", "Declaração de responsabilidade lida antes de aceitar", "Sim"),
        ("4 Identificação", "Nome do projeto, área/segmento e enquadramento conferidos", "Sim"),
        ("4 Identificação", "Tipicidade/limites de orçamento conferidos na tela de identificação", "Sim"),
        ("5 Resumo", "Resumo colado do MAPA DO PROJETO e revisado", "Sim"),
        ("6 Objetivos", "Objetivo geral e específicos coerentes com metas", "Sim"),
        ("7 Justificativa", "Justificativa específica (sem texto genérico)", "Sim"),
        ("7 Justificativa", "Metodologia / etapas de trabalho descritas", "Sim"),
        ("8 Acessibilidade", "Medidas físicas, comunicacionais e atitudinais descritas por produto", "Sim"),
        ("9 Democratização", "Medidas de democratização de acesso descritas com números", "Sim"),
        ("10 Plano de distribuição", "Cada produto cadastrado com quantidades por faixa e preços", "Sim"),
        ("11 Localização", "Local(is) de realização: país, UF, município", "Sim"),
        ("12 Cronograma", "Período de execução informado (dentro do limite de execução)", "Sim"),
        ("13 Orçamento", "Planilha orçamentária lançada por produto/etapa, com unidade, quantidade e valor unitário", "Sim"),
        ("13 Orçamento", "Custos vinculados e captação conferidos no RADAR DE LIMITES", "Sim"),
        ("14 Outras fontes", "Outras fontes informadas somente se existem e estão comprometidas", "Depende"),
        ("15 Documentação", "Todos os documentos obrigatórios anexados ANTES do envio (IN 29/2026)", "Sim"),
        ("15 Documentação", "Documentos legíveis, dentro da validade, no formato aceito", "Sim"),
        ("16 Revisão", "CHECKLIST FINAL com zero 'NÃO'", "Sim"),
        ("16 Revisão", "Lista de pendências do Salic zerada", "Sim"),
        ("17 Envio", "Proposta enviada ao MinC até 31/10 — comprovante/print salvo", "Sim"),
        ("17 Envio", "Acompanhamento de diligências no Salic e no e-mail cadastrado", "Sim"),
    ]
    for i, (p, it, ob) in enumerate(salic, start=5):
        for j, v in enumerate([p, it, ob], start=1):
            ws.cell(row=i, column=j, value=v)
        for col in range(1, 7):
            ws.cell(row=i, column=col).border = BORDER
            ws.cell(row=i, column=col).alignment = WRAP
        for col in (4, 5, 6):
            ws.cell(row=i, column=col).fill = F_INPUT
    last = 4 + len(salic)
    add_list(ws, f"D5:D{last}", ["Sim", "Não", "N/A"])
    add_list(ws, f"E5:E{last}", ["Sim", "Não", "N/A"])
    ws.cell(row=last + 2, column=2, value="Itens preparados").font = BOLD
    ws.cell(row=last + 2, column=4, value=f'=COUNTIF(D5:D{last},"Sim")&" / "&(COUNTA(B5:B{last})-COUNTIF(D5:D{last},"N/A"))')
    ws.cell(row=last + 3, column=2, value="Itens revisados").font = BOLD
    ws.cell(row=last + 3, column=4, value=f'=COUNTIF(E5:E{last},"Sim")&" / "&(COUNTA(B5:B{last})-COUNTIF(E5:E{last},"N/A"))')
    widths(ws, [22, 70, 14, 13, 13, 40])
    ws.freeze_panes = "A5"

    # ------------------------------------------------------------ CHECKLIST DOCUMENTOS
    ws = wb.create_sheet("CHECKLIST DOCUMENTOS")
    title(ws, "PLANILHA 6 — CHECKLIST DE DOCUMENTOS", "A lista oficial está no Manual do Proponente e na IN 29/2026 (e anexos). "
          "Esta planilha organiza; ela não substitui a lista oficial. 2026: documentos obrigatórios devem ser anexados na apresentação.", 7)
    header(ws, 4, ["Grupo", "Documento", "Obrigatório?", "Aplicável ao meu projeto?", "Já tenho?", "Validade", "Observação"])
    docs = [
        ("PROPONENTE", "Cadastro completo do proponente no Salic (dados pessoais ou da PJ)", "Sim", "Base para tudo."),
        ("PROPONENTE", "Documento de identificação / CPF do proponente PF ou do responsável legal", "Conforme Manual", "Confira o que o Salic pede no cadastro."),
        ("PROPONENTE", "CNPJ, ato constitutivo/estatuto/contrato social e ata de eleição da diretoria (PJ)", "Sim para PJ", "Natureza cultural deve constar no ato constitutivo/CNAE (PJ)."),
        ("PROPONENTE", "Comprovação de atuação na área cultural (portfólio no Salic)", "Depende", "Dispensa para primeiro projeto até R$ 200 mil — confira a redação oficial."),
        ("PROPONENTE", "Procuração, quando outra pessoa opera o Salic em nome do proponente", "Eventual", "Somente se houver procurador."),
        ("PROPONENTE", "Situação de regularidade (certidões) exigidas nas fases previstas", "Conforme fase", "Algumas exigências ocorrem em fases posteriores; confira."),
        ("PROJETO", "Currículo/portfólio dos principais profissionais do projeto", "Conforme Manual", "Mostra capacidade de execução."),
        ("PROJETO", "Cartas de anuência/aceite dos principais participantes", "Conforme segmento", "Artistas, curadores, instituições parceiras."),
        ("PROJETO", "Anuência/autorização do local de realização", "Conforme segmento", "Quando o local é de terceiros."),
        ("ESPECÍFICOS", "Documentos exigidos para o segmento (ex.: patrimônio, audiovisual, livro, plano anual)", "Depende do segmento", "Veja o anexo/lista do segmento na norma e no Manual."),
        ("ESPECÍFICOS", "Autorização de órgãos de patrimônio (tombamento) — projetos de patrimônio", "Depende", "Iphan/órgão estadual/municipal quando aplicável."),
        ("ESPECÍFICOS", "Autorização de direitos autorais / cessão de direitos, quando aplicável", "Depende", "Obras de terceiros."),
        ("EVENTUAIS", "Comprovante de outras fontes de recurso (cartas de compromisso)", "Eventual", "Só declare fonte que existe."),
        ("EVENTUAIS", "Orçamentos de referência / cotações para itens de maior valor", "Recomendado", "Ajuda a comprovar preço de mercado se questionado."),
        ("EVENTUAIS", "Documentos solicitados em diligência", "Eventual", "Responda dentro do prazo indicado."),
    ]
    for i, (g, d, ob, obs) in enumerate(docs, start=5):
        for j, v in enumerate([g, d, ob, None, None, None, obs], start=1):
            ws.cell(row=i, column=j, value=v)
        for col in range(1, 8):
            ws.cell(row=i, column=col).border = BORDER
            ws.cell(row=i, column=col).alignment = WRAP
        for col in (4, 5, 6):
            ws.cell(row=i, column=col).fill = F_INPUT
        ws.cell(row=i, column=6).number_format = DATE
    last = 4 + len(docs)
    for k in range(last + 1, last + 8):
        style_range(ws, k, k, 1, 7, F_INPUT)
    add_list(ws, f"D5:D{last + 7}", ["Sim", "Não"])
    add_list(ws, f"E5:E{last + 7}", ["Sim", "Não", "Em andamento"])
    ws.cell(row=last + 9, column=2, value="Documentos aplicáveis ainda pendentes").font = BOLD
    ws.cell(row=last + 9, column=3, value=f'=COUNTIFS(D5:D{last + 7},"Sim",E5:E{last + 7},"<>Sim")')
    ws.cell(row=last + 10, column=2, value="Documentos vencendo em até 30 dias").font = BOLD
    ws.cell(row=last + 10, column=3, value=f'=COUNTIFS(F5:F{last + 7},">="&TODAY(),F5:F{last + 7},"<="&TODAY()+30)')
    ws.conditional_formatting.add(f"F5:F{last + 7}", FormulaRule(formula=[f'AND(F5<>"",F5<TODAY())'], fill=F_RED))
    ws.conditional_formatting.add(f"F5:F{last + 7}", FormulaRule(formula=[f'AND(F5<>"",F5<=TODAY()+30)'], fill=F_YEL))
    widths(ws, [16, 60, 18, 16, 14, 14, 50])
    ws.freeze_panes = "A5"

    # ------------------------------------------------------------ CHECKLIST FINAL
    ws = wb.create_sheet("CHECKLIST FINAL")
    title(ws, "CHECKLIST FINAL — ANTES DE CLICAR EM ENVIAR", "Marque SIM, NÃO ou N/A. Meta: zero NÃO. Qualquer NÃO = corrija antes de enviar.", 4)
    header(ws, 4, ["Bloco", "Verificação", "SIM / NÃO / N/A", "Onde corrigir"])
    from checklist_final import ITENS
    for i, (b, t, onde) in enumerate(ITENS, start=5):
        ws.cell(row=i, column=1, value=b).font = BOLD
        ws.cell(row=i, column=2, value=t)
        ws.cell(row=i, column=3).fill = F_INPUT
        ws.cell(row=i, column=4, value=onde)
        for col in range(1, 5):
            ws.cell(row=i, column=col).border = BORDER
            ws.cell(row=i, column=col).alignment = WRAP
    last = 4 + len(ITENS)
    add_list(ws, f"C5:C{last}", ["SIM", "NÃO", "N/A"])
    ws.conditional_formatting.add(f"C5:C{last}", FormulaRule(formula=['C5="NÃO"'], fill=F_RED))
    ws.conditional_formatting.add(f"C5:C{last}", FormulaRule(formula=['C5="SIM"'], fill=F_GRN))
    s = last + 2
    for k, (lab, f) in enumerate([
        ("SIM", f'=COUNTIF(C5:C{last},"SIM")'), ("NÃO", f'=COUNTIF(C5:C{last},"NÃO")'),
        ("N/A", f'=COUNTIF(C5:C{last},"N/A")'), ("Em branco", f'=COUNTBLANK(C5:C{last})'),
    ]):
        ws.cell(row=s + k, column=2, value=lab).font = BOLD
        ws.cell(row=s + k, column=3, value=f)
    ws.cell(row=s + 4, column=2, value="RESULTADO").font = Font(bold=True, size=13)
    ws.cell(row=s + 4, column=3, value=f'=IF(C{s + 1}>0,"🔴 NÃO ENVIE AINDA: "&C{s + 1}&" item(ns) NÃO",IF(C{s + 3}>0,"🟡 Faltam "&C{s + 3}&" itens","🟢 Pronto para revisão final e envio"))')
    ws.merge_cells(start_row=s + 4, start_column=3, end_row=s + 4, end_column=4)
    traffic_light(ws, f"C{s + 4}", f"C{s + 4}")
    widths(ws, [22, 80, 16, 32])
    ws.freeze_panes = "A5"

    # ------------------------------------------------------------ CONTROLE NORMATIVO
    ws = wb.create_sheet("CONTROLE NORMATIVO")
    title(ws, "ÚLTIMA VERIFICAÇÃO NORMATIVA", "Registre cada vez que você conferir a norma oficial. Se algo mudou, atualize a aba PARAMETROS.", 5)
    header(ws, 4, ["Data", "Norma consultada", "Versão", "Alterações relevantes", "Link oficial"])
    ctrl = [
        (date(2026, 9, 25), "Instrução Normativa MinC nº 29/2026", "29/01/2026 (DOU 30/01/2026)",
         "Base deste kit. Revogou a IN MinC nº 23/2025. Novos limites por proponente, prazo máx. de execução de 36 meses, documentos obrigatórios na apresentação, acessibilidade reforçada.",
         "https://www.gov.br/cultura/pt-br/acesso-a-informacao/legislacao-e-normativas/instrucao-normativa-minc-no-29-de-29-de-janeiro-de-2026"),
        (date(2026, 9, 25), "Lei nº 8.313/1991 (Lei Rouanet)", "compilada", "Base legal do Pronac e do incentivo fiscal.", "https://www.planalto.gov.br/ccivil_03/leis/l8313cons.htm"),
        (date(2026, 9, 25), "Decreto nº 11.453/2023", "vigente", "Regulamenta os mecanismos de fomento do Pronac.", "https://www.planalto.gov.br/ccivil_03/_ato2023-2026/2023/decreto/d11453.htm"),
        (date(2026, 9, 25), "Manual do Proponente – IN 2026", "publicado pelo MinC", "Passo a passo no Salic.", "https://www.gov.br/cultura/pt-br/centrais-de-conteudo/marcas-e-logotipos/marcas-rouanet/ManualdoProponenteIN2026DEFINITIVO2.pdf"),
    ]
    for i, row in enumerate(ctrl, start=5):
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = BORDER
            c.alignment = WRAP
        ws.cell(row=i, column=1).number_format = DATE
    for k in range(5 + len(ctrl), 16):
        style_range(ws, k, k, 1, 5, F_INPUT)
    widths(ws, [14, 36, 26, 70, 70])

    # ------------------------------------------------------------ exemplo
    if example:
        fill_example(wb)

    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    fn = "ROUANET-31-10-Planilhas-EXEMPLO.xlsx" if example else "ROUANET-31-10-Planilhas.xlsx"
    path = os.path.abspath(os.path.join(OUT_DIR, fn))
    wb.save(path)
    return path


def fill_example(wb):
    """Projeto fictício para demonstração (valores ilustrativos, não são referência de preço)."""
    ws = wb["MAPA DO PROJETO"]
    mapa = [
        "Teatro na Praça: Circuito de Rua 2027 (EXEMPLO FICTÍCIO)",
        "Artes cênicas",
        "Teatro (confira o segmento exato na tabela do Salic)",
        "Circuito com 8 apresentações gratuitas de teatro de rua em 4 praças de bairros periféricos de [cidade/UF], entre março e junho de 2027, "
        "com 4 oficinas de iniciação teatral para jovens, tradução em Libras e audiodescrição em todas as sessões.",
        "8 apresentações gratuitas do espetáculo '[nome]' + 4 oficinas de iniciação teatral (12 h cada).",
        "Ampliar o acesso ao teatro em bairros sem equipamento cultural de [cidade].",
        "1) Realizar 8 apresentações gratuitas em 4 praças; 2) Formar 60 jovens em 4 oficinas; 3) Garantir acessibilidade comunicacional em 100% das sessões.",
        "Os bairros X e Y não têm teatro nem centro cultural (dados da prefeitura, 2025). O grupo atua há 6 anos na cidade...",
        "Pré-produção (mapeamento das praças, autorizações, ensaios de adaptação), produção (apresentações + oficinas), pós-produção (relatório e registro).",
        "Moradores dos bairros X e Y, famílias, crianças e jovens de 12 a 24 anos; público estimado de 2.400 pessoas.",
        "60 jovens das oficinas; escolas públicas do entorno convidadas; pessoas com deficiência (Libras e audiodescrição).",
        "8 apresentações; 4 oficinas; 1 registro audiovisual.",
        "8 apresentações; 60 vagas de oficina; 2.400 espectadores; 8 sessões com Libras e audiodescrição.",
        "[Cidade/UF] — praças dos bairros X, Y, Z e W (espaço público, com autorização da prefeitura).",
        "01/02/2027 a 31/07/2027 (6 meses).",
        "Pré-produção: fev; Produção: mar–jun; Pós-produção: jul.",
        "Rampas removíveis e área reservada; intérprete de Libras e audiodescrição em todas as sessões; equipe orientada para receber PcD.",
        "100% gratuito; ações formativas gratuitas; sessões em bairros sem equipamento cultural.",
        "2.400 pessoas atendidas; 60 jovens formados; comprovação por fotos, listas de presença, vídeo e relatório.",
        "Apoio local de lanche para as oficinas (carta de compromisso do comércio parceiro).",
    ]
    for i, v in enumerate(mapa, start=5):
        ws.cell(row=i, column=3, value=v)

    ws = wb["ORÇAMENTO"]
    INC, OUT = FONTES
    linhas = [
        ("Pré-produção", "Coordenação geral", "Coordenação do projeto pelo proponente (serviço efetivo)", "mês", 6, 4000, INC, "Proponente (PF)", "", "Remuneração do proponente", "Não", "Coordena equipe, cronograma e prestação de contas", "Sim", "Não"),
        ("Pré-produção", "Produção executiva", "Produtor(a) executivo(a)", "mês", 6, 3500, INC, "Produtora A", "", "Produção / atividade-fim", "Não", "Executa logística das 8 apresentações", "Sim", "Não"),
        ("Pré-produção", "Visitas técnicas", "Mapeamento e visita às 4 praças", "serviço", 1, 2500, INC, "Produtora A", "", "Produção / atividade-fim", "Não", "Escolha e viabilidade dos locais", "Sim", "Não"),
        ("Produção / Execução", "Cachê do grupo", "Grupo teatral – por apresentação", "apresentação", 8, 6000, INC, "Grupo Teatral B", "", "Cachê – grupo ou coletivo", "Não", "Produto principal: 8 apresentações", "Sim", "Não"),
        ("Produção / Execução", "Sonorização e luz", "Locação de som e luz – diária", "diária", 8, 1800, INC, "Som & Luz C", "", "Produção / atividade-fim", "Não", "Apresentações ao ar livre exigem som", "Sim", "Não"),
        ("Produção / Execução", "Cenografia e figurino", "Adaptação para rua", "verba", 1, 12000, INC, "Ateliê D", "", "Produção / atividade-fim", "Não", "Espetáculo adaptado ao espaço público", "Sim", "Não"),
        ("Produção / Execução", "Transporte", "Van para elenco e equipamentos", "diária", 8, 900, INC, "Transportes E", "", "Produção / atividade-fim", "Não", "Deslocamento às 4 praças", "Sim", "Não"),
        ("Produção / Execução", "Arte-educador", "Oficinas de iniciação teatral (12 h cada)", "oficina", 4, 1500, INC, "Arte-educador F", "", "Produção / atividade-fim", "Não", "Objetivo específico 2 (formar 60 jovens)", "Sim", "Não"),
        ("Produção / Execução", "Lanche oficinas", "Lanche para participantes", "oficina", 4, 300, OUT, "Comércio parceiro", "Apoio confirmado por carta", "Produção / atividade-fim", "Não", "Permanência dos jovens nas oficinas", "Sim", "Não"),
        ("Acessibilidade", "Intérprete de Libras", "Por sessão", "sessão", 8, 600, INC, "Intérprete G", "", "Acessibilidade / comunicação e divulgação acessíveis", "Não", "Acessibilidade comunicacional (objetivo 3)", "Sim", "Não"),
        ("Acessibilidade", "Audiodescrição", "Por sessão", "sessão", 8, 700, INC, "Audiodescritor H", "", "Acessibilidade / comunicação e divulgação acessíveis", "Não", "Acessibilidade comunicacional (objetivo 3)", "Sim", "Não"),
        ("Acessibilidade", "Rampas removíveis", "Locação por diária", "diária", 8, 350, INC, "Locadora I", "", "Acessibilidade / comunicação e divulgação acessíveis", "Não", "Acessibilidade física em praça", "Sim", "Não"),
        ("Acessibilidade", "Consultoria de acessibilidade", "Plano e acompanhamento", "serviço", 1, 3000, INC, "Consultor J", "", "Acessibilidade / comunicação e divulgação acessíveis", "Não", "Planejar medidas atitudinais e treinar equipe", "Sim", "Não"),
        ("Divulgação / Comunicação", "Assessoria de imprensa", "Mensal", "mês", 3, 2000, INC, "Comunicação K", "", "Divulgação / comunicação", "Não", "Levar público aos bairros", "Sim", "Não"),
        ("Divulgação / Comunicação", "Impulsionamento", "Redes sociais segmentadas por bairro", "verba", 1, 4000, INC, "Plataformas", "", "Divulgação / comunicação", "Não", "Alcançar moradores do entorno", "Sim", "Não"),
        ("Divulgação / Comunicação", "Peças gráficas", "Cartazes e filipetas (com versão acessível)", "verba", 1, 3500, INC, "Gráfica L", "", "Divulgação / comunicação", "Não", "Divulgação presencial nos bairros", "Sim", "Não"),
        ("Pós-produção", "Registro audiovisual", "Captação e edição de vídeo", "serviço", 1, 6000, INC, "Audiovisual M", "", "Produção / atividade-fim", "Não", "Produto 'registro' + comprovação", "Não", "Não"),
        ("Custos administrativos", "Contador", "Serviços contábeis", "mês", 6, 700, INC, "Contabilidade N", "", "Custos de administração", "Não", "Controle financeiro e prestação de contas", "Sim", "Não"),
        ("Custos administrativos", "Tarifas bancárias", "Conta do projeto", "mês", 6, 60, INC, "Banco", "", "Custos de administração", "Não", "Movimentação obrigatória da conta", "Sim", "Não"),
        ("Captação de recursos", "Captador", "Remuneração proporcional ao captado", "serviço", 1, 9000, INC, "Captador O", "", "Remuneração de captação", "Não", "Buscar patrocinadores", "Sim", "Não"),
    ]
    for k, row in enumerate(linhas):
        r = ORC_FIRST + k
        et, it, de, un, q, vu, fo, forn, obs, cat, vinc, porque, preco, dupla = row
        for col, v in zip([1, 2, 3, 4, 5, 6, 8, 9, 10, 11, 12, 13, 14, 15], [et, it, de, un, q, vu, fo, forn, obs, cat, vinc, porque, preco, dupla]):
            ws.cell(row=r, column=col, value=v)

    ws = wb["CRONOGRAMA"]
    ws["B3"] = date(2027, 2, 1)
    ws["B4"] = date(2027, 7, 31)
    cron = [
        ("Mapeamento das praças e autorizações", "Pré-produção", date(2027, 2, 1), date(2027, 2, 20), "Produtora A"),
        ("Plano de acessibilidade e treinamento da equipe", "Acessibilidade", date(2027, 2, 10), date(2027, 2, 28), "Consultor J"),
        ("Campanha de divulgação nos bairros", "Divulgação / Comunicação", date(2027, 2, 20), date(2027, 6, 15), "Comunicação K"),
        ("Apresentações (8 sessões)", "Produção / Execução", date(2027, 3, 6), date(2027, 6, 20), "Grupo Teatral B"),
        ("Oficinas de iniciação teatral", "Produção / Execução", date(2027, 3, 13), date(2027, 6, 13), "Arte-educador F"),
        ("Edição do registro audiovisual", "Pós-produção", date(2027, 6, 21), date(2027, 7, 15), "Audiovisual M"),
        ("Relatório final e prestação de contas", "Encerramento", date(2027, 7, 1), date(2027, 7, 31), "Proponente"),
    ]
    for k, (a, e, i, f, resp) in enumerate(cron):
        r = CR_FIRST + k
        for col, v in zip([1, 2, 3, 4, 6], [a, e, i, f, resp]):
            ws.cell(row=r, column=col, value=v)

    ws = wb["DISTRIBUICAO"]
    ws["C4"] = 2400
    for r, v in zip(range(7, 12), [0, 0, 2400, 0, 0]):
        ws.cell(row=r, column=2, value=v)
    ws["B19"] = "Sessões em bairros sem equipamento cultural + convite a escolas públicas do entorno."
    ws["B20"] = "4 oficinas gratuitas de iniciação teatral para jovens de 12 a 24 anos."

    ws = wb["RADAR DE LIMITES"]
    ws["B5"] = "Pessoa física"
    ws["B6"] = "Não"
    ws["B13"] = "Sim"
    ws["B16"] = 2460

    ws = wb["MATRIZ CONSISTENCIA"]
    mat = [
        ("Realizar 8 apresentações gratuitas", "8 apresentações / 2.400 pessoas", "Apresentação teatral", "Cachê do grupo", "Apresentações (8 sessões)", "Gratuita social – moradores X/Y", "OK"),
        ("Formar 60 jovens", "4 oficinas / 60 vagas", "Oficina", "Arte-educador", "Oficinas de iniciação teatral", "Jovens 12–24 anos", "OK"),
        ("Acessibilidade em 100% das sessões", "8 sessões com Libras e AD", "Apresentação teatral", "Intérprete de Libras", "Plano de acessibilidade e treinamento da equipe", "PcD", "OK"),
    ]
    for k, row in enumerate(mat):
        for col, v in enumerate(row, start=1):
            ws.cell(row=5 + k, column=col if col < 7 else 9, value=v)


if __name__ == "__main__":
    import sys
    sys.path.insert(0, os.path.dirname(__file__))
    os.makedirs(OUT_DIR, exist_ok=True)
    print(build(example=False))
    print(build(example=True))
